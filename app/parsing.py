"""Turn uploaded bytes (PDF / DOCX / TXT) into plain text, or fail with a precise, actionable error.

"Unreadable" is not one failure: an image-only scan, an encrypted PDF, a corrupt file and a PDF
whose fonts extract as junk all need different advice for the recruiter, so each gets its own code.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field

from docx import Document
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.schemas import ErrorCode

MAX_BYTES = 10 * 1024 * 1024


class DocumentError(Exception):
    def __init__(self, code: ErrorCode, message: str, hint: str | None = None):
        super().__init__(message)
        self.code, self.message, self.hint = code, message, hint


@dataclass
class ParsedDocument:
    text: str
    kind: str
    pages: int = 1
    warnings: list[str] = field(default_factory=list)


def _sniff(filename: str, data: bytes, allowed: set[str]) -> str:
    """Decide the type from magic bytes first, extension second, so a renamed file can't fool us."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if data.startswith(b"%PDF"):
        kind = "pdf"
    elif data.startswith(b"PK\x03\x04") and ext == "docx":
        kind = "docx"
    elif ext == "pdf":
        raise DocumentError(ErrorCode.CORRUPT_FILE, "File has a .pdf extension but is not a valid PDF.",
                            "Re-export the file as PDF, or upload it as a .txt file.")
    elif ext == "docx":
        raise DocumentError(ErrorCode.CORRUPT_FILE, "File has a .docx extension but is not a valid Word document.")
    elif ext in ("txt", "md", "text"):
        kind = "txt"
    else:
        kind = ext or "unknown"
    if kind not in allowed:
        raise DocumentError(ErrorCode.UNSUPPORTED_FILE_TYPE, f"Unsupported file type '.{ext or '?'}'.",
                            f"Supported formats: {', '.join(sorted(a.upper() for a in allowed))}.")
    return kind


def _decode_text(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf_text(data: bytes) -> tuple[str, int]:
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            # Many "encrypted" PDFs only carry an owner password and open with an empty user password.
            try:
                ok = reader.decrypt("")
            except Exception:
                ok = 0
            if not ok:
                raise DocumentError(ErrorCode.ENCRYPTED_FILE, "The PDF is password protected.",
                                    "Ask the candidate for an unprotected copy.")
        pages = [p.extract_text() or "" for p in reader.pages]
    except DocumentError:
        raise
    except (PdfReadError, ValueError, KeyError, TypeError, AttributeError, IndexError) as exc:
        raise DocumentError(ErrorCode.CORRUPT_FILE, f"The PDF could not be parsed ({type(exc).__name__}: {exc}).",
                            "The file may be damaged. Re-export it as PDF or upload a .txt/.docx version.")
    return "\n\n".join(pages), len(pages)


def _docx_text(data: bytes) -> str:
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx surfaces a variety of zip/xml errors
        raise DocumentError(ErrorCode.CORRUPT_FILE, f"The Word document could not be opened ({exc}).")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(parts)


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace(" ", " ").replace("\t", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ ​]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def garble_ratio(text: str) -> float:
    """Fraction of non-space characters that look like extraction junk.

    Broken font encodings produce "(cid:123)" runs, private-use glyphs or replacement characters
    instead of letters. A real resume is overwhelmingly letters, digits and punctuation.
    """
    stripped = re.sub(r"\s", "", text)
    if not stripped:
        return 1.0
    cid = sum(len(m) for m in re.findall(r"\(cid:\d+\)", stripped))
    junk = sum(1 for ch in stripped if ch == "�" or unicodedata.category(ch) in ("Co", "Cn", "Cc"))
    return min(1.0, (cid + junk) / len(stripped))


def extract_text(filename: str, data: bytes, *, allowed: set[str] = frozenset({"pdf", "docx", "txt"}),
                 min_chars: int = 0, max_chars: int = 60000, purpose: str = "resume") -> ParsedDocument:
    if not data:
        raise DocumentError(ErrorCode.EMPTY_FILE, f"The {purpose} file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(ErrorCode.FILE_TOO_LARGE, f"The {purpose} file exceeds {MAX_BYTES // 1024 // 1024} MB.")

    kind = _sniff(filename, data, set(allowed))
    pages = 1
    if kind == "pdf":
        raw, pages = _pdf_text(data)
    elif kind == "docx":
        raw = _docx_text(data)
    else:
        raw = _decode_text(data)
    return check_text(normalize(raw), kind=kind, pages=pages, min_chars=min_chars,
                      max_chars=max_chars, purpose=purpose)


def check_text(text: str, *, kind: str = "txt", pages: int = 1, min_chars: int = 0,
               max_chars: int = 60000, purpose: str = "resume") -> ParsedDocument:
    warnings: list[str] = []
    if kind == "pdf" and len(text) < max(50, 30 * pages):
        raise DocumentError(
            ErrorCode.NO_TEXT_LAYER,
            f"The PDF has {pages} page(s) but almost no extractable text, so it is probably a scanned image.",
            "OCR is not supported yet. Ask for a text-based PDF, or upload a .txt/.docx version.",
        )
    if not text:
        raise DocumentError(ErrorCode.INSUFFICIENT_TEXT, f"No text was found in the {purpose}.")
    ratio = garble_ratio(text)
    if ratio > 0.15:
        raise DocumentError(
            ErrorCode.GARBLED_TEXT,
            f"Text extracted from the {purpose} is unreadable ({ratio:.0%} junk characters), usually caused "
            "by embedded fonts without a Unicode map.",
            "Re-export the PDF with standard fonts, or upload a .txt/.docx version.",
        )
    if len(text) < min_chars:
        raise DocumentError(ErrorCode.INSUFFICIENT_TEXT,
                            f"Only {len(text)} characters of text were found in the {purpose}, which is too little to assess.",
                            "Check that the correct file was uploaded.")
    if len(text) > max_chars:
        warnings.append(f"{purpose.capitalize()} truncated from {len(text)} to {max_chars} characters.")
        text = text[:max_chars]
    if ratio > 0.02:
        warnings.append(f"{ratio:.1%} of extracted {purpose} characters look garbled; some content may be missing.")
    return ParsedDocument(text=text, kind=kind, pages=pages, warnings=warnings)
