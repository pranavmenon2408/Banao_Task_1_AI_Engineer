"""Convert uploaded documents (PDF, DOCX, TXT) to plain text, or fail with a precise, actionable error.

An "unreadable" document can fail in several ways: an image-only scan, an encrypted PDF, a corrupt file, or a PDF
whose fonts extract as junk. Each needs different advice for the recruiter, so each has its own error code.

PDFs are read with PyMuPDF in column-aware reading order (`pdf_layout`). A PDF without a text layer is passed to
the OCR engine (`ocr`) when one is supplied, and is rejected only if OCR fails too.
"""

from __future__ import annotations

import io
import re
import time
import unicodedata
from dataclasses import dataclass, field
from typing import TYPE_CHECKING

import pymupdf
from docx import Document

from app.core.schemas import ErrorCode, StageMetric
from app.documents import pdf_layout

if TYPE_CHECKING:
    from app.documents.ocr import OcrEngine

pymupdf.TOOLS.mupdf_display_errors(False)  # parse failures are reported as DocumentError instead

MAX_BYTES = 10 * 1024 * 1024
DEFAULT_ALLOWED = frozenset({"pdf", "docx", "txt"})


class DocumentError(Exception):
    """A document could not be turned into usable text. `hint` tells the user what to do about it."""

    def __init__(self, code: ErrorCode, message: str, hint: str | None = None) -> None:
        super().__init__(message)
        self.code, self.message, self.hint = code, message, hint


@dataclass
class ParsedDocument:
    """Extracted text plus how it was obtained."""

    text: str
    kind: str
    pages: int = 1
    warnings: list[str] = field(default_factory=list)
    method: str = ""  # "pymupdf", "pymupdf-columns", "ocr-tesseract", "ocr-vlm", "docx" or "text"
    metric: StageMetric | None = None  # extraction latency, plus vision-model tokens when OCR used one


def _sniff(filename: str, data: bytes, allowed: set[str]) -> str:
    """Decide the file type from magic bytes first and the extension second, so a renamed file is caught."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if data.startswith(b"%PDF"):
        kind = "pdf"
    elif data.startswith(b"PK\x03\x04") and ext == "docx":
        kind = "docx"
    elif ext == "pdf":
        raise DocumentError(
            ErrorCode.CORRUPT_FILE,
            "File has a .pdf extension but is not a valid PDF.",
            "Re-export the file as PDF, or upload it as a .txt file.",
        )
    elif ext == "docx":
        raise DocumentError(ErrorCode.CORRUPT_FILE, "File has a .docx extension but is not a valid Word document.")
    elif ext in ("txt", "md", "text"):
        kind = "txt"
    else:
        kind = ext or "unknown"
    if kind not in allowed:
        raise DocumentError(
            ErrorCode.UNSUPPORTED_FILE_TYPE,
            f"Unsupported file type '.{ext or '?'}'.",
            f"Supported formats: {', '.join(sorted(a.upper() for a in allowed))}.",
        )
    return kind


def _decode_text(data: bytes) -> str:
    """Decode a text file, trying UTF-16 (BOM), UTF-8 and Windows-1252 in that order."""
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _open_pdf(data: bytes) -> pymupdf.Document:
    """Open a PDF, mapping corrupt and password-protected files to DocumentError."""
    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # pymupdf raises FileDataError and several lower-level errors
        raise DocumentError(
            ErrorCode.CORRUPT_FILE,
            f"The PDF could not be parsed ({type(exc).__name__}: {exc}).",
            "The file may be damaged. Re-export it as PDF or upload a .txt/.docx version.",
        ) from exc
    # Many "encrypted" PDFs only carry an owner password and open with an empty user password.
    if doc.needs_pass and not doc.authenticate(""):
        raise DocumentError(
            ErrorCode.ENCRYPTED_FILE, "The PDF is password protected.", "Ask the candidate for an unprotected copy."
        )
    if doc.page_count == 0:
        raise DocumentError(
            ErrorCode.CORRUPT_FILE,
            "The PDF has no readable pages.",
            "The file may be damaged. Re-export it as PDF or upload a .txt/.docx version.",
        )
    return doc


def _pdf_text(data: bytes, ocr: OcrEngine | None, purpose: str, metric: StageMetric) -> tuple[str, int, str, list[str]]:
    """Return (text, pages, method, warnings) for a PDF, using OCR when it has no text layer."""
    doc = _open_pdf(data)
    try:
        raw, multi = pdf_layout.extract(doc)
    except Exception as exc:
        raise DocumentError(
            ErrorCode.CORRUPT_FILE, f"Text could not be read from the PDF ({type(exc).__name__}: {exc})."
        ) from exc
    pages = doc.page_count
    method = "pymupdf-columns" if multi else "pymupdf"
    if len(normalize(raw)) >= max(50, 30 * pages):
        return raw, pages, method, []

    # No text layer: a scan or a photo of a document.
    if ocr is None or not ocr.cfg.enabled:
        raise DocumentError(
            ErrorCode.NO_TEXT_LAYER,
            f"The PDF has {pages} page(s) but almost no extractable text, so it is probably a scanned image.",
            "OCR is disabled. Ask for a text-based PDF, or upload a .txt/.docx version.",
        )
    from app.documents.ocr import OcrFailed

    try:
        out = ocr.run(doc, metric)
    except OcrFailed as exc:
        raise DocumentError(
            ErrorCode.NO_TEXT_LAYER,
            f"The {purpose} PDF is a scanned image and OCR could not read it ({'; '.join(exc.notes)}).",
            "Ask for a text-based PDF or a clearer scan, or upload a .txt/.docx version.",
        ) from exc
    if out.method == "tesseract":
        how = "Tesseract OCR" + (f", confidence {out.confidence:.0f}/100" if out.confidence is not None else "")
    else:
        how = f"a vision model ({ocr.vlm.last_model.split('/')[-1]})"
    warnings = [
        f"The {purpose} is a scanned image; its text was read with {how}. "
        "Check quoted evidence against the original if a score looks off.",
        *out.notes,
    ]
    return out.text, pages, f"ocr-{out.method}", warnings


def _docx_text(data: bytes) -> str:
    """Paragraph and table text from a Word document."""
    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # python-docx surfaces a variety of zip/xml errors
        raise DocumentError(ErrorCode.CORRUPT_FILE, f"The Word document could not be opened ({exc}).") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
    return "\n".join(parts)


def normalize(text: str) -> str:
    """Unicode-normalise text and collapse irregular whitespace and blank lines."""
    text = unicodedata.normalize("NFKC", text)
    text = text.replace(" ", " ").replace("\t", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ ​]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def garble_ratio(text: str) -> float:
    """Fraction of non-space characters that look like extraction junk.

    Broken font encodings produce "(cid:123)" runs, private-use glyphs or replacement characters instead of
    letters, whereas real documents are overwhelmingly letters, digits and punctuation.
    """
    stripped = re.sub(r"\s", "", text)
    if not stripped:
        return 1.0
    cid = sum(len(m) for m in re.findall(r"\(cid:\d+\)", stripped))
    junk = sum(1 for ch in stripped if ch == "�" or unicodedata.category(ch) in ("Co", "Cn", "Cc"))
    return min(1.0, (cid + junk) / len(stripped))


def extract_text(
    filename: str,
    data: bytes,
    *,
    allowed: frozenset[str] = DEFAULT_ALLOWED,
    min_chars: int = 0,
    max_chars: int = 60000,
    purpose: str = "resume",
    ocr: OcrEngine | None = None,
) -> ParsedDocument:
    """Extract and validate text from an uploaded file.

    Raises:
        DocumentError: the file is empty, too large, unsupported, corrupt, encrypted or unreadable.
    """
    if not data:
        raise DocumentError(ErrorCode.EMPTY_FILE, f"The {purpose} file is empty.")
    if len(data) > MAX_BYTES:
        raise DocumentError(ErrorCode.FILE_TOO_LARGE, f"The {purpose} file exceeds {MAX_BYTES // 1024 // 1024} MB.")

    kind = _sniff(filename, data, set(allowed))
    t0 = time.perf_counter()
    metric = StageMetric(name=f"{purpose.replace(' ', '_')}_parsing", latency_ms=0)
    pages, warnings = 1, []
    if kind == "pdf":
        raw, pages, method, warnings = _pdf_text(data, ocr, purpose, metric)
    elif kind == "docx":
        raw, method = _docx_text(data), "docx"
    else:
        raw, method = _decode_text(data), "text"
    metric.latency_ms = round((time.perf_counter() - t0) * 1000, 1)
    metric.detail = method
    doc = check_text(
        normalize(raw),
        kind="ocr" if method.startswith("ocr") else kind,
        pages=pages,
        min_chars=min_chars,
        max_chars=max_chars,
        purpose=purpose,
    )
    doc.warnings = warnings + doc.warnings
    doc.method, doc.metric = method, metric
    return doc


def check_text(
    text: str,
    *,
    kind: str = "txt",
    pages: int = 1,
    min_chars: int = 0,
    max_chars: int = 60000,
    purpose: str = "resume",
) -> ParsedDocument:
    """Validate extracted text: reject empty, garbled or too-short text and truncate over-long text."""
    warnings: list[str] = []
    if kind == "pdf" and len(text) < max(50, 30 * pages):
        raise DocumentError(
            ErrorCode.NO_TEXT_LAYER,
            f"The PDF has {pages} page(s) but almost no extractable text, so it is probably a scanned image.",
            "Ask for a text-based PDF, or upload a .txt/.docx version.",
        )
    if not text:
        raise DocumentError(ErrorCode.INSUFFICIENT_TEXT, f"No text was found in the {purpose}.")
    ratio = garble_ratio(text)
    if ratio > 0.15:
        raise DocumentError(
            ErrorCode.GARBLED_TEXT,
            f"Text extracted from the {purpose} is unreadable ({ratio:.0%} junk characters), usually caused "
            "by embedded fonts without a Unicode map.",
            "Re-export the PDF with standard fonts, or upload a .txt/.docx version.",
        )
    if len(text) < min_chars:
        raise DocumentError(
            ErrorCode.INSUFFICIENT_TEXT,
            f"Only {len(text)} characters of text were found in the {purpose}, which is too little to assess.",
            "Check that the correct file was uploaded.",
        )
    if len(text) > max_chars:
        warnings.append(f"{purpose.capitalize()} truncated from {len(text)} to {max_chars} characters.")
        text = text[:max_chars]
    if ratio > 0.02:
        warnings.append(f"{ratio:.1%} of extracted {purpose} characters look garbled; some content may be missing.")
    return ParsedDocument(text=text, kind=kind, pages=pages, warnings=warnings)
