import io

import pytest
from docx import Document

from app.documents.parsing import DocumentError, extract_text, garble_ratio
from app.core.schemas import ErrorCode
from tests.conftest import make_image_only_pdf, make_text_pdf


def _code(fn):
    with pytest.raises(DocumentError) as exc:
        fn()
    return exc.value.code


def test_text_pdf_extracts(resume_lines):
    doc = extract_text("cv.pdf", make_text_pdf(resume_lines), min_chars=100)
    assert doc.kind == "pdf"
    assert "FastAPI" in doc.text and "IIT Madras" in doc.text


def test_scanned_pdf_is_rejected_as_no_text_layer():
    assert _code(lambda: extract_text("scan.pdf", make_image_only_pdf(2))) == ErrorCode.NO_TEXT_LAYER


def test_encrypted_pdf_with_user_password(resume_lines):
    data = make_text_pdf(resume_lines, encrypt="secret")
    assert _code(lambda: extract_text("locked.pdf", data)) == ErrorCode.ENCRYPTED_FILE


def test_truncated_pdf_is_corrupt_or_unreadable(resume_lines):
    data = make_text_pdf(resume_lines)[:300]
    assert _code(lambda: extract_text("broken.pdf", data)) in {ErrorCode.CORRUPT_FILE, ErrorCode.NO_TEXT_LAYER}


def test_renamed_non_pdf():
    assert _code(lambda: extract_text("cv.pdf", b"hello I am not a pdf")) == ErrorCode.CORRUPT_FILE


def test_empty_and_unsupported():
    assert _code(lambda: extract_text("cv.txt", b"")) == ErrorCode.EMPTY_FILE
    assert _code(lambda: extract_text("cv.png", b"\x89PNG....")) == ErrorCode.UNSUPPORTED_FILE_TYPE


def test_docx_not_allowed_for_resume_when_restricted():
    buf = io.BytesIO()
    Document().save(buf)
    assert _code(lambda: extract_text("cv.docx", buf.getvalue(), allowed={"pdf", "txt"})) == ErrorCode.UNSUPPORTED_FILE_TYPE


def test_docx_with_table():
    d = Document()
    d.add_paragraph("We need a Python engineer with 5+ years of experience.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Must have", "Kubernetes"
    buf = io.BytesIO()
    d.save(buf)
    doc = extract_text("jd.docx", buf.getvalue(), purpose="job description")
    assert "Kubernetes" in doc.text and "5+ years" in doc.text


def test_garbled_text_rejected():
    junk = "(cid:12)(cid:44)(cid:9)" * 40
    assert garble_ratio(junk) > 0.9
    assert _code(lambda: extract_text("cv.txt", junk.encode())) == ErrorCode.GARBLED_TEXT


def test_too_little_text():
    assert _code(lambda: extract_text("cv.txt", b"John Smith", min_chars=200)) == ErrorCode.INSUFFICIENT_TEXT


def test_utf16_and_cp1252_txt():
    assert "Café" in extract_text("a.txt", "Café résumé".encode("utf-16")).text
    assert "Café" in extract_text("a.txt", "Café résumé".encode("cp1252")).text


def test_truncation_warning():
    doc = extract_text("a.txt", ("word " * 5000).encode(), max_chars=1000)
    assert len(doc.text) == 1000 and doc.warnings
