"""Generate PDF/DOCX versions of the sample files plus deliberately unreadable resumes.

    python -m scripts.make_samples

Outputs in samples/:
  resume_a_strong.pdf              text PDF of resume A
  jd_backend_engineer.docx         the JD as a Word document
  unreadable_scanned_resume.pdf    image-only PDF (what a phone scan looks like to a parser)
  unreadable_encrypted_resume.pdf  password-protected PDF
  unreadable_corrupt_resume.pdf    truncated PDF bytes
"""
from __future__ import annotations

import io
import sys
import textwrap
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
S = ROOT / "samples"


def text_pdf(lines: list[str], encrypt: str | None = None) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4, encrypt=encrypt)
    y = 800
    for raw in lines:
        for line in textwrap.wrap(raw, 95) or [""]:
            c.setFont("Helvetica-Bold" if raw.isupper() and raw.strip() else "Helvetica", 10)
            c.drawString(50, y, line)
            y -= 14
            if y < 50:
                c.showPage()
                y = 800
    c.save()
    return buf.getvalue()


def scanned_pdf(lines: list[str]) -> bytes:
    """Render text as tiny filled rectangles: visually a page of text, but no text layer."""
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setFillGray(0.15)
    y = 800
    for line in lines:
        x = 50
        for word in line.split():
            w = 5.2 * len(word)
            c.rect(x, y, w, 7, fill=1, stroke=0)
            x += w + 4
        y -= 14
    c.showPage()
    c.save()
    return buf.getvalue()


def main() -> None:
    a = (S / "resume_a_strong.txt").read_text(encoding="utf-8").splitlines()
    (S / "resume_a_strong.pdf").write_bytes(text_pdf(a))
    (S / "unreadable_scanned_resume.pdf").write_bytes(scanned_pdf(a))
    (S / "unreadable_encrypted_resume.pdf").write_bytes(text_pdf(a, encrypt="candidate123"))
    (S / "unreadable_corrupt_resume.pdf").write_bytes(text_pdf(a)[:400])

    d = Document()
    for line in (S / "jd_backend_engineer.txt").read_text(encoding="utf-8").splitlines():
        if line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            d.add_paragraph(line)
    d.save(S / "jd_backend_engineer.docx")
    print("samples written to", S)


if __name__ == "__main__":
    sys.exit(main())
